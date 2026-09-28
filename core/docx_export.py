"""FRD → .docx through a Word template (FR-DOC-09, D-13).

The template is an ordinary .docx designed in Word — title page, approval sheet,
logo, headers and footers, fixed text — with placeholders where the app puts
content:

  inline (anywhere, incl. tables, headers, footers):
    {{title}} {{project}} {{version}} {{date}} {{today}} {{year}} {{count}} {{meta}} {{author}}
  a paragraph of its own:
    {{toc}}            table of contents
    {{body}}           every section, in the order set by the FRD skill
    {{section:KEY}}    one section (e.g. {{section:functional}}), for a custom order or text in between

Headings use the template's "Heading 1/2" styles (created if missing), so the
look is whatever the template says. Sources become real Word footnotes.
The built-in "standard" and "GOST" looks are generated as such templates, and
copying them gives an editable starting point.
"""
import datetime
import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Mm, Pt, RGBColor

from core.transcripts import format_time

LABELS = {
    "ru": {"version": "Версия", "built": "собрана", "toc": "Содержание", "conflict": "Конфликт",
           "source": "источник", "atoms": ("требование", "требования", "требований"), "empty": "Не указано."},
    "en": {"version": "Version", "built": "built", "toc": "Contents", "conflict": "Conflict",
           "source": "source", "atoms": ("requirement", "requirements", "requirements"), "empty": "Not stated."},
}
FOOTNOTES_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
PLACEHOLDER = re.compile(r"\{\{\s*([a-z_]+(?::[a-z0-9_]+)?)\s*\}\}")
INLINE = {"title", "project", "version", "date", "today", "year", "count", "meta", "author"}
PLACEHOLDER_HELP = [
    ("{{title}}", "Название документа"), ("{{project}}", "Проект"), ("{{version}}", "Номер версии"),
    ("{{date}}", "Дата сборки версии"), ("{{today}}", "Дата выгрузки"), ("{{year}}", "Год"),
    ("{{count}}", "Число требований"), ("{{meta}}", "Строка «Версия N · собрана … · N требований»"),
    ("{{author}}", "Кто собрал версию"),
    ("{{toc}}", "Содержание (отдельным абзацем)"), ("{{body}}", "Все разделы (отдельным абзацем)"),
    ("{{section:functional}}", "Один раздел по ключу из скилла сборки (отдельным абзацем)"),
]


def _plural(n, forms, lang):
    if lang != "ru":
        return forms[0] if n == 1 else forms[1]
    if n % 10 == 1 and n % 100 != 11:
        return forms[0]
    return forms[1] if 2 <= n % 10 <= 4 and not 12 <= n % 100 <= 14 else forms[2]


# ── footnotes ────────────────────────────────────────────────────────────────

class _Footnotes:
    """Adds footnotes to the template's own footnotes part, or creates one."""

    def __init__(self, doc, size):
        self.doc, self.size, self.notes = doc, size, []
        self.part = next((r.target_part for r in doc.part.rels.values() if r.reltype == RT.FOOTNOTES), None)
        self.first_id = 1
        if self.part is not None:
            ids = [int(x) for x in re.findall(rb'<w:footnote\b[^>]*\bw:id="(-?\d+)"', self.part.blob)]
            self.first_id = max([0] + ids) + 1

    def reference(self, paragraph, text):
        fid = self.first_id + len(self.notes)
        self.notes.append((fid, text))
        run = paragraph.add_run()
        rpr = OxmlElement("w:rPr")
        va = OxmlElement("w:vertAlign")
        va.set(qn("w:val"), "superscript")
        rpr.append(va)
        run._r.append(rpr)
        ref = OxmlElement("w:footnoteReference")
        ref.set(qn("w:id"), str(fid))
        run._r.append(ref)

    def _xml(self):
        hp = int(self.size.pt * 2)
        out = []
        for fid, text in self.notes:
            safe = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            out.append(f'<w:footnote {nsdecls("w")} w:id="{fid}"><w:p><w:pPr><w:spacing w:before="0" w:after="0" '
                       'w:line="240" w:lineRule="auto"/><w:ind w:firstLine="0"/><w:jc w:val="left"/></w:pPr>'
                       f'<w:r><w:rPr><w:vertAlign w:val="superscript"/><w:sz w:val="{hp}"/></w:rPr><w:footnoteRef/></w:r>'
                       f'<w:r><w:rPr><w:sz w:val="{hp}"/></w:rPr><w:t xml:space="preserve"> {safe}</w:t></w:r></w:p></w:footnote>')
        return "".join(out)

    def attach(self):
        if not self.notes:
            return
        if self.part is not None:
            blob = self.part.blob.decode("utf-8")
            # Our fragments declare the w namespace themselves; drop the duplicate declarations.
            fragment = self._xml().replace(f" {nsdecls('w')}", "")
            self.part._blob = blob.replace("</w:footnotes>", fragment + "</w:footnotes>").encode("utf-8")
            return
        xml = (f"<w:footnotes {nsdecls('w')}>"
               '<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr><w:spacing w:after="0" w:line="240" '
               'w:lineRule="auto"/></w:pPr><w:r><w:separator/></w:r></w:p></w:footnote>'
               '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr><w:spacing w:after="0" w:line="240" '
               'w:lineRule="auto"/></w:pPr><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
               + self._xml().replace(f" {nsdecls('w')}", "") + "</w:footnotes>")
        part = Part(PackURI("/word/footnotes.xml"), FOOTNOTES_CT, xml.encode("utf-8"), self.doc.part.package)
        self.doc.part.relate_to(part, RT.FOOTNOTES)
        settings = self.doc.settings.element
        if settings.find(qn("w:footnotePr")) is None:
            settings.append(parse_xml(f'<w:footnotePr {nsdecls("w")}><w:footnote w:id="-1"/><w:footnote w:id="0"/></w:footnotePr>'))


def _source_note(src, lang):
    bits = []
    if src.get("source_date"):
        bits.append(datetime.datetime.fromtimestamp(src["source_date"]).strftime("%d.%m.%Y"))
    if src.get("start") is not None:
        bits.append(format_time(src["start"]))
    who = src.get("speaker_name") or src.get("speaker")
    if who:
        bits.append({"OTHER": "—"}.get(who, who))
    head = " · ".join(bits)
    quote = f"«{src['quote']}»" if lang == "ru" else f"“{src['quote']}”"
    title = src.get("source_title") or ""
    return (f"{head} — {quote}" if head else quote) + (f" ({LABELS[lang]['source']}: {title})" if title else "")


# ── styles ───────────────────────────────────────────────────────────────────

def _set_font(style, name, size, bold=None, color=None):
    style.font.name = name
    style.font.size = size
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        fonts.attrib.pop(qn(attr), None)


def _style(doc, name):
    try:
        return doc.styles[name]
    except KeyError:
        return None


def _ensure_heading(doc, level):
    """Templates from Word may not contain Heading styles until used: create a plain one with an outline level
    (so the table of contents finds it)."""
    st = _style(doc, f"Heading {level}")
    if st is not None:
        return st
    from docx.enum.style import WD_STYLE_TYPE
    st = doc.styles.add_style(f"Heading {level}", WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = doc.styles["Normal"]
    st.font.bold = True
    st.font.size = Pt(15 if level == 1 else 12.5)
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    st.paragraph_format.space_after = Pt(6)
    ppr = st.element.get_or_add_pPr()
    lvl = OxmlElement("w:outlineLvl")
    lvl.set(qn("w:val"), str(level - 1))
    ppr.append(lvl)
    return st


# ── built-in looks, as starter templates ─────────────────────────────────────

def _placeholder(doc, text, align=None, bold=False, size=None, page_break=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
        p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.page_break_before = page_break
    r = p.add_run(text)
    r.bold = bold
    if size:
        r.font.size = size
    return p


def _page_field(paragraph):
    for kind, text in (("begin", None), ("instr", "PAGE"), ("separate", None), ("text", "1"), ("end", None)):
        run = paragraph.add_run()
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {text} "
            run._r.append(el)
        elif kind == "text":
            run.text = text
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            run._r.append(el)


def starter(kind="neutral"):
    """The built-in "standard" or "GOST" look as an editable .docx with placeholders."""
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    normal = doc.styles["Normal"]
    if kind == "gost":
        sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Mm(30), Mm(15), Mm(20), Mm(20)
        _set_font(normal, "Times New Roman", Pt(14))
        pf = normal.paragraph_format
        pf.line_spacing, pf.first_line_indent = 1.5, Cm(1.25)
        pf.space_after = pf.space_before = Pt(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for level in (1, 2):
            st = doc.styles[f"Heading {level}"]
            _set_font(st, "Times New Roman", Pt(14), bold=True, color=RGBColor(0, 0, 0))
            st.font.italic = False
            st.paragraph_format.first_line_indent = Cm(1.25)
            st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(12), Pt(6)
            st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            st.paragraph_format.keep_with_next = True
        center = WD_ALIGN_PARAGRAPH.CENTER
        for _ in range(8):
            doc.add_paragraph()
        _placeholder(doc, "ФУНКЦИОНАЛЬНЫЕ ТРЕБОВАНИЯ", center, bold=True)
        _placeholder(doc, "{{title}}", center, bold=True)
        _placeholder(doc, "{{meta}}", center)
        for _ in range(12):
            doc.add_paragraph()
        _placeholder(doc, "{{year}}", center)
        _placeholder(doc, "{{toc}}", page_break=True)
        _placeholder(doc, "{{body}}", page_break=True)
        footer = sec.footer.paragraphs[0]
        footer.alignment = center
        footer.paragraph_format.first_line_indent = Cm(0)
        _page_field(footer)
        sec.different_first_page_header_footer = True            # no page number on the title page
    else:
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = sec.bottom_margin = Cm(2)
        _set_font(normal, "Arial", Pt(10.5))
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15
        for level, size in ((1, Pt(15)), (2, Pt(12))):
            st = doc.styles[f"Heading {level}"]
            _set_font(st, "Arial", size, bold=True, color=RGBColor(0x19, 0x1C, 0x1A))
            st.paragraph_format.space_before = Pt(16 if level == 1 else 10)
            st.paragraph_format.space_after = Pt(6)
            st.paragraph_format.keep_with_next = True
        _placeholder(doc, "{{title}}", bold=True, size=Pt(20))
        meta = _placeholder(doc, "{{meta}}")
        meta.runs[0].font.color.rgb = RGBColor(0x5E, 0x64, 0x5F)
        _placeholder(doc, "{{toc}}")
        _placeholder(doc, "{{body}}", page_break=True)
        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _page_field(footer)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── rendering ────────────────────────────────────────────────────────────────

def _all_paragraphs(doc):
    """Body, tables (nested) and every header/footer."""
    def from_container(container):
        for p in container.paragraphs:
            yield p
        for table in getattr(container, "tables", []):
            for row in table.rows:
                for cell in row.cells:
                    yield from from_container(cell)
    yield from from_container(doc)
    for section in doc.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer,
                     section.even_page_header, section.even_page_footer):
            if part is not None and not part.is_linked_to_previous:
                yield from from_container(part)


def _replace_inline(paragraph, values):
    text = "".join(r.text for r in paragraph.runs)
    if "{{" not in text:
        return
    new = PLACEHOLDER.sub(lambda m: values.get(m.group(1), m.group(0)) if m.group(1) in INLINE else m.group(0), text)
    if new == text:
        return
    for run in paragraph.runs:                       # placeholder inside one run: keep its formatting
        if "{{" in run.text:
            replaced = PLACEHOLDER.sub(lambda m: values.get(m.group(1), m.group(0)) if m.group(1) in INLINE else m.group(0), run.text)
            if run.text != replaced:
                run.text = replaced
    if "".join(r.text for r in paragraph.runs) != new:  # split across runs (Word does that): merge into the first
        paragraph.runs[0].text = new
        for run in paragraph.runs[1:]:
            run.text = ""


class _Writer:
    """Inserts content before an anchor paragraph (the placeholder), or at the end of the body."""

    def __init__(self, doc, anchor, suffix, notes, lang):
        self.doc, self.anchor, self.suffix, self.notes, self.lang = doc, anchor, suffix, notes, lang
        self.first = True
        self.page_break = bool(anchor is not None and anchor.paragraph_format.page_break_before)

    def para(self, text="", style=None):
        if self.anchor is not None:
            p = self.anchor.insert_paragraph_before(text)
        else:
            p = self.doc.add_paragraph(text)
        if style is not None:
            p.style = style
        if self.first:
            self.first = False
            if self.page_break:
                p.paragraph_format.page_break_before = True       # the placeholder asked for a new page
        return p

    def heading(self, number, title, level):
        return self.para(f"{number}{self.suffix} {title}", _ensure_heading(self.doc, level))

    def bullet(self, text):
        st = _style(self.doc, "List Bullet")
        return self.para(text, st) if st is not None else self.para(f"– {text}")

    def req(self, block):
        p = self.para()
        p.add_run(f"{block['id']}. ").bold = True
        p.add_run(block["text"])
        for src in block.get("sources") or []:
            self.notes.reference(p, _source_note(src, self.lang))
        if block.get("conflict"):
            r = self.para().add_run(f"⚠ {LABELS[self.lang]['conflict']}: {block['conflict']}")
            r.italic = True
            r.font.color.rgb = RGBColor(0x9E, 0x36, 0x26)

    def blocks(self, blocks):
        for b in blocks:
            if b["kind"] == "req":
                self.req(b)
            elif b["kind"] == "text":
                for line in (x.strip() for x in b["text"].split("\n")):
                    if line:
                        self.para(line)
            elif b["kind"] == "list":
                if b.get("title"):
                    self.para().add_run(b["title"]).bold = True
                for item in b["items"]:
                    self.bullet(item)

    def section(self, sec, free):
        self.heading(sec["number"], sec["title"], 1)
        for text in free.get(sec["key"], []):
            self.para(text)
        self.blocks(sec["blocks"])
        for sub in sec.get("subsections") or []:
            self.heading(sub["number"], sub["title"], 2)
            self.blocks(sub["blocks"])
        if not sec["blocks"] and not sec.get("subsections") and not free.get(sec["key"]):
            self.para(LABELS[self.lang]["empty"])

    def toc(self, content):
        self.para().add_run(LABELS[self.lang]["toc"]).bold = True
        entries = []
        for sec in content["sections"]:
            entries.append((1, f"{sec['number']}{self.suffix} {sec['title']}"))
            for sub in sec.get("subsections") or []:
                entries.append((2, f"{sub['number']}{self.suffix} {sub['title']}"))
        paragraphs = []
        for level, text in entries:
            st = _style(self.doc, f"TOC {level}")
            p = self.para(style=st)
            if st is None:
                p.paragraph_format.left_indent = Cm(0.8 * (level - 1))
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.space_after = Pt(0)
            paragraphs.append((p, text))
        if not paragraphs:
            return
        # One TOC field around pre-filled entries: readable at once, Word adds page numbers on update.
        first, last = paragraphs[0][0], paragraphs[-1][0]
        for kind in ("begin", "instr", "separate"):
            run = first.add_run()
            if kind == "instr":
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = ' TOC \\o "1-2" \\h \\z \\u '
            else:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
                if kind == "begin":
                    el.set(qn("w:dirty"), "true")
            run._r.append(el)
        for p, text in paragraphs:
            p.add_run(text)
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        last.add_run()._r.append(end)
        settings = self.doc.settings.element
        if settings.find(qn("w:updateFields")) is None:
            upd = OxmlElement("w:updateFields")
            upd.set(qn("w:val"), "true")
            settings.append(upd)


def _values(document, version, project_name, lang):
    l = LABELS[lang]
    n = version["atom_count"]
    built = datetime.datetime.fromtimestamp(version["created_at"])
    return {"title": document["title"], "project": project_name or "", "version": str(version["number"]),
            "date": built.strftime("%d.%m.%Y"), "today": datetime.date.today().strftime("%d.%m.%Y"),
            "year": str(datetime.date.today().year), "count": str(n), "author": version.get("created_by") or "",
            "meta": f"{l['version']} {version['number']} · {l['built']} {built.strftime('%d.%m.%Y %H:%M')} · "
                    f"{n} {_plural(n, l['atoms'], lang)}"}


def render(document, version, free_blocks, template=None, numbering="dot", project_name=""):
    """Bytes of the .docx: the template (bytes of a .docx; the standard look if None) filled with the version."""
    content = version["content"]
    lang = content.get("language", "ru")
    doc = Document(io.BytesIO(template if template else starter("neutral")))
    suffix = "." if numbering == "dot" else ""
    body_size = _style(doc, "Normal").font.size
    notes = _Footnotes(doc, Pt(10) if body_size is not None and body_size.pt >= 13 else Pt(8))
    values = _values(document, version, project_name, lang)
    for p in list(_all_paragraphs(doc)):
        _replace_inline(p, values)
    free = {}
    for fb in free_blocks:
        free.setdefault(fb["section"], []).append(fb["text"])
    by_key = {s["key"]: s for s in content["sections"]}
    placed_body = False
    for p in list(doc.paragraphs):
        m = PLACEHOLDER.fullmatch(p.text.strip())
        if not m or m.group(1) in INLINE:
            continue
        name = m.group(1)
        w = _Writer(doc, p, suffix, notes, lang)
        if name == "toc":
            w.toc(content)
        elif name == "body":
            for sec in content["sections"]:
                w.section(sec, free)
            placed_body = True
        elif name.startswith("section:") and name.split(":", 1)[1] in by_key:
            w.section(by_key[name.split(":", 1)[1]], free)
            placed_body = True
        p._p.getparent().remove(p._p)
    if not placed_body:                              # a template without {{body}}: add everything at the end
        w = _Writer(doc, None, suffix, notes, lang)
        for sec in content["sections"]:
            w.section(sec, free)
    notes.attach()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
