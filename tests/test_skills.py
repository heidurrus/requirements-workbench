import io
import zipfile

import pytest

from core import docx_export, frd, skills
from core.skills import SkillError
from core.store import Store
from tests.test_frd import PREFS, full_reply, llm, seed
from tests.test_recorder import wait_for


@pytest.fixture(autouse=True)
def data_dir(tmp_path, monkeypatch):
    monkeypatch.setenv("WORKBENCH_DATA_DIR", str(tmp_path / "data"))


@pytest.fixture()
def store(tmp_path):
    return Store(root=str(tmp_path / "lib"), user="ba")


# ── the built-in set ─────────────────────────────────────────────────────────

def test_every_builtin_skill_loads_and_every_stage_has_a_default():
    loaded = skills.all_skills()
    assert all(s.error is None for s in loaded), [(s.name, s.error) for s in loaded if s.error]
    resolved = skills.resolve()
    assert {st: s.name for st, s in resolved.items()} == skills.DEFAULTS


def test_compose_adds_contract_and_house_rules():
    base = skills.resolve()
    house = skills.create("house-rules", title="Наши правила")
    skills.save(house.name, instructions="Всегда пиши «Заказчик», а не «клиент».")
    skills.set_global("global", house.name)
    text = skills.compose(skills.resolve(), "extract", "CONTRACT-X")
    assert text.startswith(base["extract"].instructions[:40])
    assert "## Format (set by the app, always follow)\nCONTRACT-X" in text
    assert text.endswith("Всегда пиши «Заказчик», а не «клиент».")


# ── editing your own ─────────────────────────────────────────────────────────

def test_copy_edit_history_restore():
    s = skills.create("extract-requirements", title="Жёсткое извлечение")
    assert not s.builtin and s.name == "extract-requirements-copy" and s.title == "Жёсткое извлечение"
    skills.save(s.name, instructions="Только требования с числами.")
    skills.save(s.name, instructions="Только требования с числами и ролями.", description="строже")
    now = skills.get(s.name)
    assert now.instructions == "Только требования с числами и ролями." and now.version == 3
    hist = skills.history(s.name)
    assert len(hist) == 2 and all(h["kind"] == "text" for h in hist)
    first = hist[-1]["id"]
    assert "You are a senior business analyst" in skills.history_text(s.name, first)
    skills.restore(s.name, first)
    assert skills.get(s.name).instructions.startswith("You are a senior business analyst")
    assert len(skills.history(s.name)) == 3                       # restoring keeps the replaced version too


def test_builtin_skills_are_read_only():
    with pytest.raises(SkillError, match="copy"):
        skills.save("write-frd", instructions="x")
    with pytest.raises(SkillError, match="copy"):
        skills.delete("write-frd")


def test_validation_keeps_requirements_placeable():
    s = skills.create("write-frd")
    sections = [x for x in s.meta["sections"] if x["key"] != "questions"]
    with pytest.raises(SkillError, match="questions"):
        skills.save(s.name, settings={"sections": sections})
    with pytest.raises(SkillError, match="instructions"):
        skills.save(s.name, settings={"sections": s.meta["sections"] + [{"key": "glossary", "title": "Глоссарий"}]})
    with pytest.raises(SkillError, match="empty"):
        skills.save(s.name, instructions="   ")
    assert skills.get(s.name).meta["sections"] == s.meta["sections"], "a failed save changes nothing"


def test_delete_and_undelete():
    s = skills.create("fix-requirement")
    skills.delete(s.name)
    assert s.name not in {x.name for x in skills.all_skills()}
    skills.undelete(s.name)
    assert skills.get(s.name).stage == "fix"


def test_broken_skill_is_listed_with_its_error_and_others_still_work():
    import os
    bad = os.path.join(skills.custom_dir(), "broken-one")
    os.makedirs(bad)
    with open(os.path.join(bad, "SKILL.md"), "w", encoding="utf-8") as f:
        f.write("no header here")
    listed = {s.name: s for s in skills.all_skills()}
    assert "header" in listed["broken-one"].error and listed["write-frd"].error is None


def test_a_broken_or_missing_choice_falls_back_to_the_default():
    assert skills.resolve({"frd": "does-not-exist", "summary": "write-frd"})["frd"].name == "write-frd"
    assert skills.resolve({"summary": "write-frd"})["summary"].name == "summarize-source"   # wrong stage


# ── sharing ──────────────────────────────────────────────────────────────────

def test_export_import_roundtrip_renames_on_clash():
    s = skills.create("quality-check", title="Банковские правила")
    skills.save(s.name, settings={"rules": [{"id": "method", "title": "Метод измерения",
                                             "description": "Каждое NFR называет способ измерения."}]})
    data = skills.export_zip(s.name)
    imported = skills.import_file("share.zip", data)
    assert imported.name == f"{s.name}-2" and imported.meta["rules"][0]["id"] == "method"
    names = zipfile.ZipFile(io.BytesIO(data)).namelist()
    assert f"{s.name}/SKILL.md" in names and not any(".history" in n for n in names)


def test_import_single_skill_md_and_reject_bad_archives():
    md = b"---\nname: my-summary\ntitle: My summary\nstage: summary\n---\nSummarise in three bullets.\n"
    assert skills.import_file("SKILL.md", md).instructions == "Summarise in three bullets."
    evil = io.BytesIO()
    with zipfile.ZipFile(evil, "w") as z:
        z.writestr("x/SKILL.md", md.decode())
        z.writestr("x/../../escape.md", "boom")
    with pytest.raises(SkillError, match="unsafe"):
        skills.import_file("evil.zip", evil.getvalue())
    with pytest.raises(SkillError, match="SKILL.md"):
        skills.import_file("empty.zip", _zip({"readme.txt": "hi"}))
    with pytest.raises(SkillError):
        skills.import_file("x.zip", b"not a zip")


def test_import_skips_scripts():
    md = "---\nname: with-script\ntitle: T\nstage: fix\n---\nFix it.\n"
    s = skills.import_file("s.zip", _zip({"with-script/SKILL.md": md, "with-script/run.sh": "rm -rf /"}))
    import os
    assert not os.path.exists(os.path.join(s.path, "run.sh"))


def _zip(files):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for name, text in files.items():
            z.writestr(name, text)
    return buf.getvalue()


# ── skills drive the pipeline ────────────────────────────────────────────────

def test_frd_skill_controls_sections_and_extra_ai_sections(store):
    pid, *_ = seed(store)
    s = skills.create("write-frd", title="С глоссарием")
    sections = [{"key": "questions", "title": "Сначала вопросы"}] + [x for x in s.meta["sections"] if x["key"] != "questions"]
    sections.append({"key": "glossary", "title": "Глоссарий", "instructions": "Термины из требований с определениями."})
    skills.save(s.name, settings={"sections": sections}, instructions="Пиши коротко. Язык: {language}.")
    reply = dict(full_reply(), extra=[{"key": "glossary", "text": "Карточка — экран клиента."}])
    fake = llm(reply)
    frd.build(store, pid, PREFS, "k", "", complete=fake, skillset=skills.resolve({"frd": s.name}))
    system = fake.calls[0]["system"]
    assert system.startswith("Пиши коротко. Язык: Russian.") and "glossary (“Глоссарий”)" in system
    v = store.version(store.document(pid)["id"])
    assert [x["title"] for x in v["content"]["sections"]][0] == "Сначала вопросы"
    assert v["content"]["sections"][-1]["blocks"][0]["text"] == "Карточка — экран клиента."
    assert v["content"]["skills"]["frd"] == s.name


def test_quality_skill_custom_rules_and_vague_words(store):
    pid, *_ = seed(store)
    q = skills.create("quality-check")
    skills.save(q.name, settings={"vague_words": {"ru": ["маскировать"], "en": []},
                                  "rules": [{"id": "role", "title": "Нет роли", "description": "Требование не называет роль."}]})
    reply = dict(full_reply(), issues=[{"id": "FR-1", "rule": "role", "message": "Кто открывает карточку?"}])
    fake = llm(reply)
    frd.build(store, pid, PREFS, "k", "", complete=fake, skillset=skills.resolve({"quality": q.name}))
    assert "role" in fake.calls[0]["schema"]["properties"]["issues"]["items"]["properties"]["rule"]["enum"]
    v = store.version(store.document(pid)["id"])
    by = {b["id"]: b for _n, _k, b in frd.req_blocks(v["content"])}
    assert by["FR-1"]["issues"] == [{"rule": "role", "message": "Кто открывает карточку?"}]
    assert {i["rule"] for i in by["FR-2"]["issues"]} == {"vague"}          # "маскировать" is vague here
    assert "быстро" in by["NFR-1"]["text"] and "vague" not in {i["rule"] for i in by["NFR-1"]["issues"]}
    assert v["content"]["rule_titles"] == {"role": "Нет роли"}


def test_extract_uses_the_chosen_skill(store):
    from core import atoms
    pid, _ids, sid = seed(store, accept=False)
    s = skills.create("extract-requirements")
    skills.save(s.name, instructions="EXTRACT ONLY SECURITY THINGS")
    fake = llm({"atoms": []}, {"duplicates": [], "conflicts": []})
    atoms.extract_atoms(store, sid, PREFS, "k", "", complete=fake, skillset=skills.resolve({"extract": s.name}))
    assert fake.calls[0]["system"].startswith("EXTRACT ONLY SECURITY THINGS") and "[S…]" in fake.calls[0]["system"]


# ── Word templates ───────────────────────────────────────────────────────────

def _built(store):
    pid, *_ = seed(store)
    frd.build(store, pid, PREFS, "k", "", complete=llm(full_reply()))
    doc = store.document(pid)
    return doc, store.version(doc["id"])


def test_custom_word_template_placeholders_everywhere(store):
    import docx
    doc_meta, version = _built(store)
    t = docx.Document()
    t.sections[0].header.paragraphs[0].text = "ООО «Ромашка» · {{title}}"
    p = t.add_paragraph()
    p.add_run("Документ: {{ti")
    p.add_run("tle}}, версия {{version}}")                    # Word splits placeholders across runs
    table = t.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Согласовано"
    table.cell(0, 1).text = "{{today}} · {{author}}"
    t.add_paragraph("{{section:questions}}")
    t.add_paragraph("Постоянный текст шаблона между разделами.")
    t.add_paragraph("{{section:functional}}")
    buf = io.BytesIO()
    t.save(buf)
    out = docx.Document(io.BytesIO(docx_export.render(doc_meta, version, [], template=buf.getvalue(), project_name="P")))
    texts = [x.text for x in out.paragraphs]
    assert "Документ: FRD — My project, версия 1" in texts
    assert out.sections[0].header.paragraphs[0].text == "ООО «Ромашка» · FRD — My project"
    assert out.tables[0].cell(0, 1).text.endswith("· ba")
    # numbers follow the FRD skill (and the table of contents), even when the template reorders sections
    i_q, i_fixed, i_fr = texts.index("6. Открытые вопросы"), texts.index("Постоянный текст шаблона между разделами."), \
        texts.index("3. Функциональные требования")
    assert i_q < i_fixed < i_fr and not any("{{" in x for x in texts)


def test_template_without_body_placeholder_gets_content_at_the_end(store):
    import docx
    doc_meta, version = _built(store)
    t = docx.Document()
    t.add_paragraph("Титул")
    buf = io.BytesIO()
    t.save(buf)
    out = docx.Document(io.BytesIO(docx_export.render(doc_meta, version, [], template=buf.getvalue())))
    texts = [x.text for x in out.paragraphs]
    assert texts[0] == "Титул" and "1. Назначение документа" in texts


def test_template_with_its_own_footnotes_part_is_extended(store):
    doc_meta, version = _built(store)
    first = docx_export.render(doc_meta, version, [], template=docx_export.starter("neutral"))
    second = docx_export.render(doc_meta, version, [], template=first)       # already has footnotes 1..4
    notes = zipfile.ZipFile(io.BytesIO(second)).read("word/footnotes.xml").decode()
    import re
    assert sorted(int(x) for x in re.findall(r'<w:footnote [^>]*w:id="(\d+)"', notes))[-1] == 8


def test_copying_a_builtin_export_skill_gives_an_editable_template(store):
    import os
    s = skills.create("export-gost", title="ГОСТ + наш логотип")
    assert s.meta["template"] == "template.docx" and os.path.isfile(os.path.join(s.path, "template.docx"))
    body = zipfile.ZipFile(io.BytesIO(skills.template_bytes(s))).read("word/document.xml").decode()
    assert "{{body}}" in body and "{{toc}}" in body
    new = docx_export.starter("neutral")
    skills.set_template(s.name, new)
    assert skills.template_bytes(skills.get(s.name)) == new
    assert any(h["kind"] == "template" for h in skills.history(s.name))
    with pytest.raises(SkillError, match="Word"):
        skills.set_template(s.name, b"not a docx")


# ── API ──────────────────────────────────────────────────────────────────────

@pytest.fixture()
def lib(app_module, monkeypatch, tmp_path):
    s = Store(root=str(tmp_path / "lib"), user="ba")
    monkeypatch.setattr(app_module, "library", s)
    monkeypatch.setattr(app_module.settings, "load_settings", lambda: dict(PREFS))
    monkeypatch.setattr(app_module.settings, "secret", lambda name: "key")
    return s


def test_api_list_copy_save_choose_and_project_override(client, lib):
    pid = lib.current_project()["id"]
    body = client.get("/api/skills").get_json()
    frd_stage = next(s for s in body["stages"] if s["id"] == "frd")
    assert frd_stage == {"id": "frd", "default": "write-frd", "global": "write-frd", "project": None, "effective": "write-frd"}
    one = client.get("/api/skills/write-frd").get_json()
    assert one["builtin"] and "sections" in one["meta"] and "exactly one item" in one["contract"]

    copy = client.post("/api/skills", json={"from": "write-frd", "title": "Мой FRD"}).get_json()
    r = client.put(f"/api/skills/{copy['name']}", json={"instructions": "Пиши как юрист. {language}"})
    assert r.status_code == 200 and r.get_json()["version"] == 2 and len(r.get_json()["history"]) == 1
    assert client.put("/api/skills/write-frd", json={"instructions": "x"}).status_code == 400

    client.put("/api/skills/active", json={"stage": "frd", "skill": "write-frd-gost"})
    r = client.put("/api/skills/active", json={"stage": "frd", "skill": copy["name"], "project_id": pid}).get_json()
    frd_stage = next(s for s in r["stages"] if s["id"] == "frd")
    assert (frd_stage["global"], frd_stage["project"], frd_stage["effective"]) == ("write-frd-gost", copy["name"], copy["name"])
    assert client.put("/api/skills/active", json={"stage": "summary", "skill": "write-frd"}).status_code == 400
    r = client.put("/api/skills/active", json={"stage": "frd", "skill": None, "project_id": pid}).get_json()
    assert next(s for s in r["stages"] if s["id"] == "frd")["effective"] == "write-frd-gost"


def test_api_export_import_and_template(client, lib):
    r = client.get("/api/skills/export-gost/export.zip")
    assert r.status_code == 200
    names = zipfile.ZipFile(io.BytesIO(r.data)).namelist()
    assert "export-gost/template.docx" in names                    # the built-in look is shared as a real file
    imported = client.post("/api/skills/import", data={"file": (io.BytesIO(r.data), "gost.zip")},
                           content_type="multipart/form-data").get_json()
    assert imported["name"] == "export-gost-2" and imported["meta"]["template"] == "template.docx"
    t = client.get(f"/api/skills/{imported['name']}/template.docx")
    assert t.status_code == 200 and t.data[:2] == b"PK"
    up = client.post(f"/api/skills/{imported['name']}/template",
                     data={"file": (io.BytesIO(docx_export.starter("neutral")), "t.docx")}, content_type="multipart/form-data")
    assert up.status_code == 200 and up.get_json()["has_template_file"]
    assert client.post("/api/skills/import", data={}, content_type="multipart/form-data").status_code == 400


def test_api_export_uses_the_chosen_export_skill(client, lib):
    pid, *_ = seed(lib)
    frd.build(lib, pid, PREFS, "k", "", complete=llm(full_reply()))
    doc = lib.document(pid)
    import docx
    t = docx.Document()
    t.add_paragraph("НАШ БЛАНК {{project}}")
    t.add_paragraph("{{body}}")
    buf = io.BytesIO()
    t.save(buf)
    s = skills.create("export-standard", title="Бланк")
    skills.set_template(s.name, buf.getvalue())
    r = client.get(f"/api/documents/{doc['id']}/export.docx?template={s.name}")
    assert r.status_code == 200
    first = docx.Document(io.BytesIO(r.data)).paragraphs[0].text
    assert first == "НАШ БЛАНК My project"
    assert client.get(f"/api/documents/{doc['id']}/export.docx?template=write-frd").status_code == 400


def test_api_try_a_draft_without_saving(client, app_module, lib, monkeypatch):
    pid, ids, sid = seed(lib, accept=False)
    fake = llm({"atoms": [{"type": "nfr", "statement": "Карточка открывается быстро",
                           "evidence": [{"segment": 1, "quote": "открывается быстро"}]}]})
    from core import atoms
    real = atoms.extract_candidates
    monkeypatch.setattr(atoms, "extract_candidates", lambda *a, **k: real(*a, **k, complete=fake))
    job = client.post("/api/skills/try", json={"name": "extract-requirements", "instructions": "DRAFT RULES",
                                               "source_id": sid}).get_json()["job_id"]
    assert wait_for(lambda: client.get(f"/job/{job}").get_json()["status"] == "done")
    result = client.get(f"/job/{job}").get_json()["result"]
    assert result["kind"] == "atoms" and result["atoms"][0]["statement"] == "Карточка открывается быстро"
    assert fake.calls[0]["system"].startswith("DRAFT RULES")
    assert skills.get("extract-requirements").instructions != "DRAFT RULES", "trying never saves"
    assert len(lib.list_atoms(pid)) == 4, "trying adds no atoms"
    assert client.post("/api/skills/try", json={"name": "extract-requirements"}).status_code == 400   # no source
